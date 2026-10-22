"""Turn the assessment into files you can hand over: a submission, a report, CSVs, a workbook, a
backup. Also the submit-and-lock step, which records a hash of what was submitted."""

import csv
import hashlib
import html
import io
import re
import zipfile
from datetime import datetime, timezone
from pathlib import Path

import yaml

from grc_lite import __version__
from grc_lite.models import rating
from grc_lite.web import files, records

SUBMISSION = "submission.md"
SUBMITTED = "submitted.yaml"
FORMULA_START = ("=", "+", "-", "@", "\t", "\r")


def is_locked(assessment: Path) -> bool:
    return (assessment / SUBMITTED).exists()


def _cell(value) -> str:
    return str(value if value is not None else "").replace("|", "\\|").replace("\n", " ").strip()


def _table(headers: list[str], rows: list[list]) -> list[str]:
    lines = ["| " + " | ".join(headers) + " |", "| " + " | ".join("---" for _ in headers) + " |"]
    lines += ["| " + " | ".join(_cell(c) for c in row) + " |" for row in rows]
    return lines + [""]


def _join(value) -> str:
    return ", ".join(value) if isinstance(value, list) else str(value or "")


# ---- the submission file: the notes-tool format, so `grc import-notes` reads it back ----


def _finding_block(f: dict, criteria: dict[str, dict], scale: int) -> list[str]:
    linked = [
        f"{cid} ({criteria[cid].get('source', '')}): {criteria[cid].get('rule', '')}"
        if cid in criteria
        else cid
        for cid in f.get("criteria") or []
    ]
    likelihood, impact = f.get("likelihood"), f.get("impact")
    rated = ""
    if isinstance(likelihood, int) and isinstance(impact, int):
        try:
            rated = f" = {likelihood * impact} {rating(likelihood, impact, scale)}"
        except StopIteration:
            pass
    lines = [
        f"### {f['id']}: {_cell(f.get('title')) or '(untitled)'}",
        f"- **Criteria:** {_cell('; '.join(linked))}",
        f"- **Condition:** {_cell(f.get('condition'))}",
    ]
    if f.get("cause"):
        lines.append(f"- **Cause:** {_cell(f['cause'])}")
    lines += [
        f"- **Consequence:** {_cell(f.get('consequence'))}",
        (
            f"- **Action:** {_cell(f.get('action'))} (Owner: {_cell(f.get('owner'))}; "
            f"Due: {_cell(f.get('due'))})"
        ),
        (
            f"- **Rating:** Likelihood {likelihood or '?'} ({_cell(f.get('likelihood_reason'))}) x "
            f"Impact {impact or '?'} ({_cell(f.get('impact_reason'))}){rated} [scale 1–{scale}]"
        ),
    ]
    if f.get("control_ref"):
        lines.append(f"- **Control ref:** {_cell(f['control_ref'])}")
    return lines + [""]


def submission_md(level: Path, assessment: Path) -> str:
    """Everything logged, as one Markdown file. It has no timestamps, so the same records always
    give the same file and the same hash."""
    meta = records.read_meta(assessment)
    data = {kind: records.read_records(assessment, kind) for kind in records.KINDS}
    criteria = {c["id"]: c for c in data["criteria"] if "id" in c}
    out = [f"# {meta.name}", ""]
    for path in files.list_files(level):
        target = files.notes_path(assessment, path)
        if target.exists() and target.read_text().strip():
            out += [f"## {path}", "", target.read_text().strip(), ""]
    out += ["## Criteria sheet", ""]
    out += _table(
        ["ID", "Source", "Rule", "Evidence needed", "Ev # received", "Result"],
        [
            [c.get("id"), c.get("source"), c.get("rule"), c.get("evidence_needed"),
             _join(c.get("evidence")), c.get("result", "not tested")]
            for c in data["criteria"]
        ],
    )  # fmt: skip
    out += ["## Evidence log", ""]
    out += _table(
        ["ID", "Date", "Source", "Type", "What it shows", "Reliability note"],
        [[e.get(k) for k in ("id", "date", "source", "type", "shows", "reliability")]
         for e in data["evidence"]],
    )  # fmt: skip
    out += ["## Candidate findings", ""]
    out += _table(
        ["ID", "Title", "Status", "Evidence strength", "Why"],
        [[f.get(k) for k in ("id", "title", "status", "strength", "strength_reason")]
         for f in data["findings"]],
    )  # fmt: skip
    selected = [f for f in data["findings"] if f.get("status") == "selected"]
    if selected:
        out += ["## Findings", ""]
        for f in selected:
            out += _finding_block(f, criteria, meta.scale)
    out += ["## Evidence requests", ""]
    out += _table(
        ["ID", "Asked", "What", "Asked who", "Received", "Evidence", "Notes"],
        [[r.get(k) for k in ("id", "asked", "what", "asked_who", "received", "evidence", "notes")]
         for r in data["requests"]],
    )  # fmt: skip
    out += ["## Risk register", ""]
    out += _table(
        ["ID", "Risk", "Owner", "Likelihood", "Impact", "Treatment", "Status", "Links"],
        [
            [k.get("id"), k.get("title"), k.get("owner"), k.get("likelihood"), k.get("impact"),
             k.get("treatment"), k.get("status", "open"),
             _join([*(k.get("findings") or []), *(k.get("criteria") or [])])]
            for k in data["risks"]
        ],
    )  # fmt: skip
    return "\n".join(out).rstrip() + "\n"


# ---- submit and lock ----


def _sha(text: str) -> str:
    return hashlib.sha256(text.encode()).hexdigest()


def submit(level: Path, assessment: Path) -> dict:
    """Write the submission, record its hash and lock the assessment. The lock is a commitment
    to yourself, not a security boundary: delete submitted.yaml to reopen the level."""
    text = submission_md(level, assessment)
    status = {
        "submitted_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "sha256": _sha(text),
        "problems": len(records.problems(assessment)),
        "app": f"grc-lite {__version__}",
    }
    assessment.mkdir(parents=True, exist_ok=True)
    temp = assessment / (SUBMISSION + ".tmp")
    temp.write_text(text)
    temp.replace(assessment / SUBMISSION)  # no half-written file; the status is written last
    records._write(assessment / SUBMITTED, status)
    return {**status, "markdown": text, "file": str(assessment / SUBMISSION)}


def submission_status(level: Path, assessment: Path) -> dict:
    path = assessment / SUBMITTED
    if not path.exists():
        return {"submitted": False}
    recorded = yaml.safe_load(path.read_text()) or {}
    written = assessment / SUBMISSION
    return {
        "submitted": True,
        "submitted_at": recorded.get("submitted_at"),
        "sha256": recorded.get("sha256"),
        "problems": recorded.get("problems"),
        "file": str(written),
        "file_intact": written.exists() and _sha(written.read_text()) == recorded.get("sha256"),
        "matches_records": _sha(submission_md(level, assessment)) == recorded.get("sha256"),
    }


# ---- csv and xlsx ----


def neutralise(value):
    """A spreadsheet runs a cell that starts with = + - or @ as a formula, so a note could carry
    one. Prefixing an apostrophe makes it plain text."""
    if isinstance(value, str) and value.startswith(FORMULA_START):
        return "'" + value
    return value


def _rows(assessment: Path, kind: str) -> tuple[list[str], list[list]]:
    model, _ = records.KINDS[kind]
    headers = list(model.model_fields)
    rows = [
        [_join(r[h]) if isinstance(r.get(h), list) else r.get(h, "") for h in headers]
        for r in records.read_records(assessment, kind)
    ]
    return headers, rows


def csv_bytes(assessment: Path, kind: str) -> bytes:
    headers, rows = _rows(assessment, kind)
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(headers)
    for row in rows:
        writer.writerow([neutralise(str(v) if v is not None else "") for v in row])
    return buffer.getvalue().encode("utf-8-sig")  # the BOM lets Excel read accents correctly


def xlsx_bytes(assessment: Path) -> bytes:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError:
        raise records.RecordError(
            "Excel export needs openpyxl: pip install 'grc-lite[xlsx]'"
        ) from None
    book = Workbook()
    book.remove(book.active)
    for kind in records.KINDS:
        sheet = book.create_sheet(kind.capitalize())
        headers, rows = _rows(assessment, kind)
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for row in rows:
            sheet.append([str(v) if v not in (None, "") else None for v in row])
        for line in sheet.iter_rows(min_row=2):
            for cell in line:
                if isinstance(cell.value, str) and cell.value.startswith(FORMULA_START):
                    cell.value = "'" + cell.value  # stored as text, never evaluated
                    cell.data_type = "s"
        sheet.freeze_panes = "A2"
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


# ---- backup and printable report ----


def zip_bytes(assessment: Path) -> bytes:
    if not assessment.is_dir():
        raise records.RecordError("nothing to back up yet")
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for path in sorted(assessment.rglob("*")):
            if path.is_file() and not path.is_symlink() and path.suffix != ".tmp":
                archive.write(path, path.relative_to(assessment).as_posix())
    return buffer.getvalue()


def _inline(text: str) -> str:
    text = html.escape(text)
    text = re.sub(r"\*\*([^*]+)\*\*", r"<strong>\1</strong>", text)
    return re.sub(r"`([^`]+)`", r"<code>\1</code>", text)


def _cells(line: str) -> list[str]:
    parts = re.split(r"(?<!\\)\|", line.strip().strip("|"))
    return [p.strip().replace("\\|", "|") for p in parts]


def markdown_to_html(text: str) -> str:
    """Render the report's Markdown (headings, tables, lists, paragraphs). Everything is escaped."""
    out: list[str] = []
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if re.match(r"#{1,6} ", line):
            level = len(line.split(" ")[0])
            out.append(f"<h{level}>{_inline(line[level + 1 :])}</h{level}>")
        elif (
            line.startswith("|")
            and i + 1 < len(lines)
            and re.match(r"^\|[\s:|-]+\|$", lines[i + 1])
        ):
            out.append("<table><thead><tr>" + "".join(f"<th>{_inline(c)}</th>" for c in _cells(line))
                       + "</tr></thead><tbody>")  # fmt: skip
            i += 2
            while i < len(lines) and lines[i].startswith("|"):
                out.append(
                    "<tr>" + "".join(f"<td>{_inline(c)}</td>" for c in _cells(lines[i])) + "</tr>"
                )
                i += 1
            out.append("</tbody></table>")
            continue
        elif line.startswith("- "):
            out.append("<ul>")
            while i < len(lines) and lines[i].startswith("- "):
                out.append(f"<li>{_inline(lines[i][2:])}</li>")
                i += 1
            out.append("</ul>")
            continue
        elif line.strip():
            out.append(f"<p>{_inline(line)}</p>")
        i += 1
    return "\n".join(out)


PAGE = """<!doctype html>
<html lang="en"><head><meta charset="utf-8">
<meta http-equiv="Content-Security-Policy" content="default-src 'none'; style-src 'unsafe-inline'">
<title>{title}</title>
<style>
body {{ font: 14px/1.5 system-ui, sans-serif; max-width: 900px; margin: 2rem auto; padding: 0 1rem; color: #111; }}
table {{ border-collapse: collapse; width: 100%; margin: 0 0 1rem; }}
th, td {{ border: 1px solid #bbb; padding: 4px 8px; text-align: left; vertical-align: top; }}
th {{ background: #eee; }}
h1, h2, h3 {{ page-break-after: avoid; }} tr {{ page-break-inside: avoid; }}
code {{ background: #eee; padding: 0 3px; }}
@media print {{ body {{ margin: 0; max-width: none; }} }}
</style></head><body>
{body}
</body></html>
"""


def report_html(title: str, markdown: str) -> bytes:
    return PAGE.format(title=html.escape(title), body=markdown_to_html(markdown)).encode()
