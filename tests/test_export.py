import io
import shutil
import zipfile
from pathlib import Path

import pytest
import yaml
from fastapi.testclient import TestClient

from grc_lite.importer import import_notes
from grc_lite.web import export
from grc_lite.web.app import create_app

ROOT = Path(__file__).parent.parent
FIXTURE = ROOT / "tests" / "fixtures" / "notes.md"
TOKEN = "test-token-123"
AUTH = {"X-Grc-Token": TOKEN}
JSON = {"X-Grc-Token": TOKEN, "Content-Type": "application/json"}


@pytest.fixture
def level(tmp_path):
    root = tmp_path / "level"
    root.mkdir()
    (root / "01-scope.md").write_text("# Scope\n")
    (root / "policy.md").write_text("# Policy\n")
    return root


@pytest.fixture
def work(tmp_path):
    target = tmp_path / "assessment"
    shutil.copytree(ROOT / "data" / "example", target)  # a complete, valid assessment
    return target


@pytest.fixture
def client(level, work):
    return TestClient(create_app(level, work, token=TOKEN), base_url="http://127.0.0.1:8765")


def test_submission_is_deterministic_and_in_the_notes_tool_format(level, work):
    (work / "notes").mkdir()
    (work / "notes" / "policy.md.md").write_text("my note on the policy")
    first = export.submission_md(level, work)
    assert first == export.submission_md(level, work)
    assert "## policy.md\n\nmy note on the policy" in first
    for heading in ("## Criteria sheet", "## Evidence log", "## Candidate findings", "## Findings"):
        assert heading in first


def test_submission_round_trips_through_the_importer(level, work):
    files, scale, warnings = import_notes(export.submission_md(level, work))
    assert warnings == []
    assert scale == 3
    assert [c["id"] for c in files["criteria"]] == ["C-01", "C-02", "C-03"]
    assert files["criteria"][0]["evidence"] == ["E-01", "E-02"]
    assert files["criteria"][0]["result"] == "gap"
    assert [e["id"] for e in files["evidence"]] == ["E-01", "E-02", "E-03"]
    finding = files["findings"][0]
    assert finding["criteria"] == ["C-01"]
    assert (finding["likelihood"], finding["impact"]) == (2, 3)
    assert finding["owner"] == "IT Operations Manager" and finding["due"] == "2026-06-30"
    assert finding["control_ref"] == "AC-2"
    assert [f["status"] for f in files["findings"]] == ["selected", "candidate", "dropped"]


def test_submit_locks_everything_and_records_a_matching_hash(client, work):
    assert client.get("/api/submission", headers=AUTH).json() == {"submitted": False}
    done = client.post("/api/submit", json={}, headers=JSON).json()
    assert (work / "submission.md").read_text() == done["markdown"]
    status = client.get("/api/submission", headers=AUTH).json()
    assert status["sha256"] == done["sha256"]
    assert status["file_intact"] and status["matches_records"]
    assert client.get("/api/records", headers=AUTH).json()["locked"] is True
    for call in (
        lambda: client.put("/api/notes", json={"path": "policy.md", "text": "x"}, headers=JSON),
        lambda: client.post("/api/records/criteria", json={"fields": {}}, headers=JSON),
        lambda: client.put("/api/records/criteria/C-01", json={"fields": {}}, headers=JSON),
        lambda: client.delete("/api/records/criteria/C-01", headers=JSON),
        lambda: client.put("/api/meta", json={"name": "x", "scale": 3, "slots": 8}, headers=JSON),
        lambda: client.post("/api/import-notes", json={"text": ""}, headers=JSON),
        lambda: client.post("/api/submit", json={}, headers=JSON),
    ):
        assert call().status_code == 423
    assert client.get("/api/records", headers=AUTH).status_code == 200  # reading still works


def test_editing_the_files_after_submitting_is_detected(client, work):
    client.post("/api/submit", json={}, headers=JSON)
    criteria = yaml.safe_load((work / "criteria.yaml").read_text())
    criteria[0]["result"] = "meets"
    (work / "criteria.yaml").write_text(yaml.safe_dump(criteria))
    status = client.get("/api/submission", headers=AUTH).json()
    assert status["file_intact"] is True and status["matches_records"] is False
    (work / "submission.md").write_text("changed")
    assert client.get("/api/submission", headers=AUTH).json()["file_intact"] is False


def test_deleting_submitted_yaml_reopens_the_level(client, work):
    client.post("/api/submit", json={}, headers=JSON)
    (work / "submitted.yaml").unlink()
    body = {"path": "policy.md", "text": "x"}
    assert client.put("/api/notes", json=body, headers=JSON).status_code == 200


def test_unknown_export_is_404_and_downloads_are_attachments(client):
    assert client.get("/api/export/passwd", headers=AUTH).status_code == 404
    assert client.get("/api/export/..%2Fsecret", headers=AUTH).status_code == 404
    assert client.get("/api/export/criteria.csv").status_code == 401
    got = client.get("/api/export/criteria.csv", headers=AUTH)
    assert got.headers["content-disposition"] == 'attachment; filename="criteria.csv"'
    assert got.content.startswith(b"\xef\xbb\xbfid,source,rule")


def test_csv_neutralises_spreadsheet_formulas(client, work):
    evil = '=HYPERLINK("http://evil.example","click")'
    client.post(
        "/api/records/evidence",
        json={"fields": {"shows": evil, "source": "+cmd|calc"}},
        headers=JSON,
    )
    text = client.get("/api/export/evidence.csv", headers=AUTH).content.decode("utf-8-sig")
    assert "'=HYPERLINK" in text and "'+cmd" in text
    assert ",=HYPERLINK" not in text and ',"=HYPERLINK' not in text


def test_xlsx_stores_formulas_as_text(client):
    from openpyxl import load_workbook

    evil = "=1+1"
    client.post("/api/records/evidence", json={"fields": {"shows": evil}}, headers=JSON)
    book = load_workbook(io.BytesIO(client.get("/api/export/records.xlsx", headers=AUTH).content))
    cells = [c for row in book["Evidence"].iter_rows(min_row=2) for c in row if c.value]
    hit = next(c for c in cells if "1+1" in str(c.value))
    assert hit.data_type == "s" and hit.value == "'=1+1"
    assert book.sheetnames == ["Criteria", "Evidence", "Findings", "Requests", "Risks"]


def test_zip_backup_has_the_records_and_nothing_stray(client, work, tmp_path):
    (work / "notes").mkdir()
    (work / "notes" / "a.md.md").write_text("note")
    (work / "half.tmp").write_text("partial")
    (tmp_path / "outside.txt").write_text("OUTSIDE")
    (work / "link.txt").symlink_to(tmp_path / "outside.txt")
    names = zipfile.ZipFile(
        io.BytesIO(client.get("/api/export/records.zip", headers=AUTH).content)
    ).namelist()
    assert "criteria.yaml" in names and "notes/a.md.md" in names
    assert "half.tmp" not in names and "link.txt" not in names


def test_report_html_escapes_everything_and_blocks_scripts(client):
    client.put("/api/records/criteria/C-01", json={"fields": {
        "source": "<script>alert(1)</script>", "rule": '<img src=x onerror="alert(2)">'}}, headers=JSON)  # fmt: skip
    page = client.get("/api/export/report.html", headers=AUTH).text
    assert "<script>alert" not in page and "<img src=x" not in page
    assert "&lt;script&gt;" in page
    assert "default-src 'none'" in page


def test_report_export_waits_for_a_valid_assessment(client):
    client.post("/api/records/criteria", json={"fields": {}}, headers=JSON)  # blank, so invalid
    assert client.get("/api/export/report.md", headers=AUTH).status_code == 422
    assert client.get("/api/export/report.html", headers=AUTH).status_code == 422


def test_import_fills_an_empty_assessment_and_brings_file_notes_across(tmp_path, level):
    empty = tmp_path / "empty"
    client = TestClient(create_app(level, empty, token=TOKEN), base_url="http://127.0.0.1:8765")
    text = FIXTURE.read_text()
    done = client.post(
        "/api/import-notes", json={"text": text, "name": "Demo"}, headers=JSON
    ).json()
    assert done["warnings"] == [] and done["meta"]["name"] == "Demo"
    assert done["meta"]["counters"] == {"C": 3, "E": 3, "F": 3}
    assert done["notes_imported"] == 1
    assert "Northwind" in (empty / "notes" / "01-scope.md.md").read_text()
    assert [c["id"] for c in done["criteria"]] == ["C-01", "C-02", "C-03"]
    again = client.post("/api/import-notes", json={"text": text}, headers=JSON)
    assert again.status_code == 409
    assert (
        client.post("/api/records/criteria", json={"fields": {}}, headers=JSON).json()["record"][
            "id"
        ]
        == "C-04"
    )


def test_import_refuses_huge_text(tmp_path, level):
    client = TestClient(
        create_app(level, tmp_path / "e", token=TOKEN), base_url="http://127.0.0.1:8765"
    )
    big = {"text": "x" * 2_000_001}
    assert client.post("/api/import-notes", json=big, headers=JSON).status_code == 422
